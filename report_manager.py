"""
Financial report module.
Computes total building income (charge payments) vs total expenses
(general expenses + water bills paid by the building), and can export
the full report to an Excel file.
"""

import os
from datetime import datetime

import ui
import payment_manager
import expense_manager
import water_manager


def build_summary():
    total_income = payment_manager.get_total_payments()
    total_general_expenses = expense_manager.get_total_expenses()
    total_water_expenses = water_manager.get_total_water_bills()
    total_expenses = total_general_expenses + total_water_expenses
    net_balance = total_income - total_expenses

    return {
        "total_income": total_income,
        "total_general_expenses": total_general_expenses,
        "total_water_expenses": total_water_expenses,
        "total_expenses": total_expenses,
        "net_balance": net_balance,
    }


def show_report():
    ui.print_header("Financial Report")
    summary = build_summary()

    ui.render_table(
        "Building Financial Summary",
        ["Item", "Amount"],
        [
            ("Total income (charge deposits)", ui.format_currency(summary["total_income"])),
            ("Total general expenses", ui.format_currency(summary["total_general_expenses"])),
            ("Total water bill expenses", ui.format_currency(summary["total_water_expenses"])),
            ("Total expenses (all)", ui.format_currency(summary["total_expenses"])),
            ("Net balance", ui.format_currency(summary["net_balance"])),
        ],
    )

    balance = summary["net_balance"]
    if balance >= 0:
        ui.print_success(f"Building balance is positive: {ui.format_currency(balance)}")
    else:
        ui.print_error(f"Building balance is negative: {ui.format_currency(balance)}")

    if ui.ask_confirm("\nExport this report to an Excel file?", default=False):
        path = export_to_excel(summary)
        ui.print_success(f"Report exported to: {path}")

    ui.pause()


def export_to_excel(summary: dict) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    header_fill = PatternFill(start_color="1F6FEB", end_color="1F6FEB", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Building Financial Report"
    ws["A1"].font = title_font
    ws["A2"] = f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

    headers = ["Item", "Amount"]
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = header_font
        cell.fill = header_fill

    rows = [
        ("Total income (charge deposits)", summary["total_income"]),
        ("Total general expenses", summary["total_general_expenses"]),
        ("Total water bill expenses", summary["total_water_expenses"]),
        ("Total expenses (all)", summary["total_expenses"]),
        ("Net balance", summary["net_balance"]),
    ]
    for row in rows:
        ws.append(row)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    # --- Payments sheet ---
    ws2 = wb.create_sheet("Payments")
    ws2.append(["Unit", "Amount", "Date"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for p in payment_manager.get_all_payments():
        ws2.append([p["unit_number"], p["amount"], p["payment_date"]])
    for col, width in zip("ABC", (12, 16, 14)):
        ws2.column_dimensions[col].width = width

    # --- Expenses sheet ---
    ws3 = wb.create_sheet("Expenses")
    ws3.append(["Title", "Amount", "Date"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    for e in expense_manager.get_all_expenses():
        ws3.append([e["title"], e["amount"], e["expense_date"]])
    for col, width in zip("ABC", (28, 16, 14)):
        ws3.column_dimensions[col].width = width

    # --- Water bills sheet ---
    ws4 = wb.create_sheet("Water Bills")
    ws4.append(["Month", "Total Amount", "Total People", "Cost per Person"])
    for cell in ws4[1]:
        cell.font = header_font
        cell.fill = header_fill
    for b in water_manager.get_all_water_bills():
        ws4.append([b["bill_month"], b["total_amount"], b["total_people"], b["cost_per_person"]])
    for col, width in zip("ABCD", (14, 16, 14, 16)):
        ws4.column_dimensions[col].width = width

    output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")
    os.makedirs(output_dir, exist_ok=True)
    filename = f"building_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    return filepath
