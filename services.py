"""
Service layer: all database read/write logic for the Building Management
web app. Kept free of any Flask/HTTP concerns so it's easy to test and reuse.
Dates are handled in the Jalali (Shamsi) calendar, stored as 'YYYY-MM-DD'
strings (which still sort correctly with plain SQL ORDER BY).
"""

import os
from datetime import datetime

import jdatetime

from database import get_connection

EXPORT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")


# --------------------------------------------------------------------------
# Date helpers
# --------------------------------------------------------------------------

def today_jalali() -> str:
    return jdatetime.date.today().strftime("%Y-%m-%d")


def is_valid_jalali_date(value: str) -> bool:
    try:
        jdatetime.datetime.strptime(value, "%Y-%m-%d")
        return True
    except (ValueError, TypeError):
        return False


# --------------------------------------------------------------------------
# Settings & units
# --------------------------------------------------------------------------

def get_settings():
    conn = get_connection()
    row = conn.execute("SELECT * FROM settings WHERE id = 1").fetchone()
    conn.close()
    return row


def get_units():
    conn = get_connection()
    rows = conn.execute("SELECT * FROM units ORDER BY unit_number").fetchall()
    conn.close()
    return rows


def get_unit(unit_number: int):
    conn = get_connection()
    row = conn.execute(
        "SELECT * FROM units WHERE unit_number = ?", (unit_number,)
    ).fetchone()
    conn.close()
    return row


def get_unit_numbers():
    return [u["unit_number"] for u in get_units()]


def complete_setup(total_units: int, charge_amount: float, units_people: list):
    """units_people: list of (unit_number, people_count) tuples."""
    conn = get_connection()
    cur = conn.cursor()
    cur.execute(
        """INSERT INTO settings (id, total_units, charge_amount, is_configured)
           VALUES (1, ?, ?, 1)""",
        (total_units, charge_amount),
    )
    cur.executemany(
        "INSERT INTO units (unit_number, people_count) VALUES (?, ?)",
        units_people,
    )
    conn.commit()
    conn.close()


def update_charge_amount(amount: float):
    conn = get_connection()
    conn.execute(
        "UPDATE settings SET charge_amount = ?, updated_at = CURRENT_TIMESTAMP WHERE id = 1",
        (amount,),
    )
    conn.commit()
    conn.close()


def update_unit_people(unit_number: int, people_count: int):
    conn = get_connection()
    conn.execute(
        "UPDATE units SET people_count = ? WHERE unit_number = ?",
        (people_count, unit_number),
    )
    conn.commit()
    conn.close()


def add_unit(unit_number: int, people_count: int):
    conn = get_connection()
    existing = [r["unit_number"] for r in conn.execute("SELECT unit_number FROM units")]
    if unit_number in existing:
        conn.close()
        return False
    conn.execute(
        "INSERT INTO units (unit_number, people_count) VALUES (?, ?)",
        (unit_number, people_count),
    )
    conn.execute(
        "UPDATE settings SET total_units = total_units + 1, updated_at = CURRENT_TIMESTAMP WHERE id = 1"
    )
    conn.commit()
    conn.close()
    return True


def delete_unit(unit_number: int):
    conn = get_connection()
    conn.execute("DELETE FROM units WHERE unit_number = ?", (unit_number,))
    conn.execute(
        "UPDATE settings SET total_units = total_units - 1, updated_at = CURRENT_TIMESTAMP WHERE id = 1"
    )
    conn.commit()
    conn.close()


# --------------------------------------------------------------------------
# Payments
# --------------------------------------------------------------------------

def list_payments():
    conn = get_connection()
    rows = conn.execute(
        "SELECT * FROM payments ORDER BY payment_date DESC, id DESC"
    ).fetchall()
    conn.close()
    return rows


def get_payment(payment_id: int):
    conn = get_connection()
    row = conn.execute("SELECT * FROM payments WHERE id = ?", (payment_id,)).fetchone()
    conn.close()
    return row


def add_payment(unit_number: int, amount: float, payment_date: str):
    conn = get_connection()
    conn.execute(
        "INSERT INTO payments (unit_number, amount, payment_date) VALUES (?, ?, ?)",
        (unit_number, amount, payment_date),
    )
    conn.commit()
    conn.close()


def update_payment(payment_id: int, unit_number: int, amount: float, payment_date: str):
    conn = get_connection()
    conn.execute(
        "UPDATE payments SET unit_number = ?, amount = ?, payment_date = ? WHERE id = ?",
        (unit_number, amount, payment_date, payment_id),
    )
    conn.commit()
    conn.close()


def delete_payment(payment_id: int):
    conn = get_connection()
    conn.execute("DELETE FROM payments WHERE id = ?", (payment_id,))
    conn.commit()
    conn.close()


def get_total_payments() -> float:
    conn = get_connection()
    row = conn.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM payments").fetchone()
    conn.close()
    return row["total"]


# --------------------------------------------------------------------------
# Water bills
# --------------------------------------------------------------------------

def _recompute_shares(conn, water_bill_id: int, total_amount: float):
    """Delete and re-insert shares for a water bill based on current units."""
    units = conn.execute("SELECT * FROM units ORDER BY unit_number").fetchall()
    total_people = sum(u["people_count"] for u in units)
    cost_per_person = (total_amount / total_people) if total_people > 0 else 0

    conn.execute("DELETE FROM water_bill_shares WHERE water_bill_id = ?", (water_bill_id,))
    for u in units:
        share = round(cost_per_person * u["people_count"], 2)
        conn.execute(
            """INSERT INTO water_bill_shares
               (water_bill_id, unit_number, people_count, share_amount)
               VALUES (?, ?, ?, ?)""",
            (water_bill_id, u["unit_number"], u["people_count"], share),
        )
    conn.execute(
        "UPDATE water_bills SET total_people = ?, cost_per_person = ? WHERE id = ?",
        (total_people, cost_per_person, water_bill_id),
    )
    return total_people, cost_per_person


def calculate_water_bill(bill_month: str, total_amount: float):
    """Create a new water bill and split it across units by people count."""
    conn = get_connection()
    units = conn.execute("SELECT * FROM units ORDER BY unit_number").fetchall()
    total_people = sum(u["people_count"] for u in units)
    if total_people <= 0:
        conn.close()
        return None
    cost_per_person = total_amount / total_people

    cur = conn.cursor()
    cur.execute(
        """INSERT INTO water_bills (bill_month, total_amount, total_people, cost_per_person)
           VALUES (?, ?, ?, ?)""",
        (bill_month, total_amount, total_people, cost_per_person),
    )
    water_bill_id = cur.lastrowid
    for u in units:
        share = round(cost_per_person * u["people_count"], 2)
        cur.execute(
            """INSERT INTO water_bill_shares
               (water_bill_id, unit_number, people_count, share_amount)
               VALUES (?, ?, ?, ?)""",
            (water_bill_id, u["unit_number"], u["people_count"], share),
        )
    conn.commit()
    conn.close()
    return water_bill_id


def update_water_bill(water_bill_id: int, bill_month: str, total_amount: float):
    conn = get_connection()
    conn.execute(
        "UPDATE water_bills SET bill_month = ?, total_amount = ? WHERE id = ?",
        (bill_month, total_amount, water_bill_id),
    )
    _recompute_shares(conn, water_bill_id, total_amount)
    conn.commit()
    conn.close()


def delete_water_bill(water_bill_id: int):
    conn = get_connection()
    conn.execute("DELETE FROM water_bill_shares WHERE water_bill_id = ?", (water_bill_id,))
    conn.execute("DELETE FROM water_bills WHERE id = ?", (water_bill_id,))
    conn.commit()
    conn.close()


def list_water_bills():
    conn = get_connection()
    rows = conn.execute("SELECT * FROM water_bills ORDER BY id DESC").fetchall()
    conn.close()
    return rows


def get_water_bill(water_bill_id: int):
    conn = get_connection()
    row = conn.execute("SELECT * FROM water_bills WHERE id = ?", (water_bill_id,)).fetchone()
    conn.close()
    return row


def get_water_bill_shares(water_bill_id: int):
    conn = get_connection()
    rows = conn.execute(
        "SELECT * FROM water_bill_shares WHERE water_bill_id = ? ORDER BY unit_number",
        (water_bill_id,),
    ).fetchall()
    conn.close()
    return rows


def get_total_water_bills() -> float:
    conn = get_connection()
    row = conn.execute("SELECT COALESCE(SUM(total_amount), 0) AS total FROM water_bills").fetchone()
    conn.close()
    return row["total"]


# --------------------------------------------------------------------------
# General expenses
# --------------------------------------------------------------------------

def list_expenses():
    conn = get_connection()
    rows = conn.execute(
        "SELECT * FROM expenses ORDER BY expense_date DESC, id DESC"
    ).fetchall()
    conn.close()
    return rows


def get_expense(expense_id: int):
    conn = get_connection()
    row = conn.execute("SELECT * FROM expenses WHERE id = ?", (expense_id,)).fetchone()
    conn.close()
    return row


def add_expense(title: str, amount: float, expense_date: str):
    conn = get_connection()
    conn.execute(
        "INSERT INTO expenses (title, amount, expense_date) VALUES (?, ?, ?)",
        (title, amount, expense_date),
    )
    conn.commit()
    conn.close()


def update_expense(expense_id: int, title: str, amount: float, expense_date: str):
    conn = get_connection()
    conn.execute(
        "UPDATE expenses SET title = ?, amount = ?, expense_date = ? WHERE id = ?",
        (title, amount, expense_date, expense_id),
    )
    conn.commit()
    conn.close()


def delete_expense(expense_id: int):
    conn = get_connection()
    conn.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
    conn.commit()
    conn.close()


def get_total_expenses() -> float:
    conn = get_connection()
    row = conn.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM expenses").fetchone()
    conn.close()
    return row["total"]


# --------------------------------------------------------------------------
# Financial report
# --------------------------------------------------------------------------

def build_summary():
    total_income = get_total_payments()
    total_general_expenses = get_total_expenses()
    total_water_expenses = get_total_water_bills()
    total_expenses = total_general_expenses + total_water_expenses
    net_balance = total_income - total_expenses
    return {
        "total_income": total_income,
        "total_general_expenses": total_general_expenses,
        "total_water_expenses": total_water_expenses,
        "total_expenses": total_expenses,
        "net_balance": net_balance,
    }


def export_to_excel() -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    summary = build_summary()
    wb = Workbook()

    header_fill = PatternFill(start_color="1F6FEB", end_color="1F6FEB", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Building Financial Report"
    ws["A1"].font = title_font
    ws["A2"] = f"Generated at: {today_jalali()}"

    ws.append([])
    ws.append(["Item", "Amount"])
    for cell in ws[4]:
        cell.font = header_font
        cell.fill = header_fill
    for row in [
        ("Total income (charge deposits)", summary["total_income"]),
        ("Total general expenses", summary["total_general_expenses"]),
        ("Total water bill expenses", summary["total_water_expenses"]),
        ("Total expenses (all)", summary["total_expenses"]),
        ("Net balance", summary["net_balance"]),
    ]:
        ws.append(row)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    ws2 = wb.create_sheet("Payments")
    ws2.append(["Unit", "Amount", "Date"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for p in list_payments():
        ws2.append([p["unit_number"], p["amount"], p["payment_date"]])
    for col, width in zip("ABC", (12, 16, 14)):
        ws2.column_dimensions[col].width = width

    ws3 = wb.create_sheet("Expenses")
    ws3.append(["Title", "Amount", "Date"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    for e in list_expenses():
        ws3.append([e["title"], e["amount"], e["expense_date"]])
    for col, width in zip("ABC", (28, 16, 14)):
        ws3.column_dimensions[col].width = width

    ws4 = wb.create_sheet("Water Bills")
    ws4.append(["Month", "Total Amount", "Total People", "Cost per Person"])
    for cell in ws4[1]:
        cell.font = header_font
        cell.fill = header_fill
    for b in list_water_bills():
        ws4.append([b["bill_month"], b["total_amount"], b["total_people"], b["cost_per_person"]])
    for col, width in zip("ABCD", (14, 16, 14, 16)):
        ws4.column_dimensions[col].width = width

    os.makedirs(EXPORT_DIR, exist_ok=True)
    filename = f"building_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)
    return filepath
